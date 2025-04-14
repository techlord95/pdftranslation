from flask import Flask, request, send_file, Response, jsonify
from werkzeug.utils import secure_filename
import os
import json
from advanced_docx_translator import translate_docx_advanced, translation_cache, summarize_text_groq
# from pdf_to_word import convert_pdf_exactly
# from pdf_to_word import smart_pdf_to_word
from pdf_to_word import convert_pdf_with_aspose
import logging
import time
import threading
from docx import Document
import uuid
import shutil

app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = 'uploads'
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024  # 16MB max file size

# Ensure upload directory exists
os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)

# Configure logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

class TranslationProgress:
    def __init__(self):
        self.progress = 0
        self.status = "Ready"
        self.api_calls = 0
        self.cached = 0
        self.messages = []
        self.summary = None

    def update(self, progress=None, status=None, api_calls=None, cached=None, message=None, summary=None):
        if progress is not None:
            self.progress = progress
        if status is not None:
            self.status = status
        if api_calls is not None:
            self.api_calls = api_calls
        if cached is not None:
            self.cached = cached
        if message is not None:
            self.messages.append(message)
            logger.info(message)
        if summary is not None:
            self.summary = summary

    def to_json(self):
        data = {
            'progress': self.progress,
            'status': self.status,
            'apiCalls': self.api_calls,
            'cached': self.cached,
            'message': self.messages[-1] if self.messages else None
        }
        if self.summary is not None:
            data['summary'] = self.summary
        return json.dumps(data)

@app.route('/')
def index():
    return send_file('web_translator.html')

@app.route('/styles.css')
def styles():
    return send_file('styles.css')

@app.route('/translator.js')
def translator_js():
    return send_file('translator.js')

@app.route('/translate', methods=['POST'])
def translate():
    if 'file' not in request.files:
        return 'No file provided', 400
    
    file = request.files['file']
    if file.filename == '':
        return 'No file selected', 400
    
    if not (file.filename.lower().endswith('.docx') or file.filename.lower().endswith('.pdf')):
        return 'Invalid file type. Please upload a DOCX or PDF file.', 400
    
    target_lang = request.form.get('targetLang', 'es')
    engine = request.form.get('engine', 'google')
    generate_summary = request.form.get('generateSummary') == 'true'
    # Get the first_page_only preference from the form
    first_page_only = request.form.get('firstPageOnly') == 'true' 
    
    # Save uploaded file
    original_filename = secure_filename(file.filename)
    input_path = os.path.join(app.config['UPLOAD_FOLDER'], original_filename)
    file.save(input_path)
    logger.info(f"File saved to {input_path}")
    
    current_filename = original_filename # Keep track of the filename for output
    docx_file_to_process = input_path    # Path to the DOCX file (original or converted)

    # If the file is a PDF, convert it to DOCX using Aspose
    if original_filename.lower().endswith('.pdf'):
        pdf_path = input_path
        docx_filename = f"{os.path.splitext(original_filename)[0]}.docx"
        docx_path = os.path.join(app.config['UPLOAD_FOLDER'], docx_filename)
        logger.info(f"Detected PDF. Attempting conversion to: {docx_path}")
        try:
            convert_pdf_with_aspose(pdf_path, docx_path)
            docx_file_to_process = docx_path
            current_filename = docx_filename 
            logger.info(f"PDF successfully converted to DOCX using Aspose: {docx_path}")
        except Exception as e:
            logger.error(f"Aspose PDF conversion failed: {str(e)}")
            return f"PDF conversion failed: {str(e)}", 500
        finally:
            # Optionally remove the original PDF 
            if os.path.exists(pdf_path) and docx_file_to_process == docx_path:
                try:
                    os.remove(pdf_path)
                    logger.info(f"Removed original PDF: {pdf_path}")
                except OSError as rm_err:
                    logger.error(f"Error removing original PDF {pdf_path}: {rm_err}")
            pass 
    
    # Define output paths based on the (potentially converted) filename
    output_path = os.path.join(app.config['UPLOAD_FOLDER'], f'translated_{os.path.splitext(current_filename)[0]}.docx')
    download_path = os.path.join(app.config['UPLOAD_FOLDER'], f'download_{os.path.splitext(current_filename)[0]}.docx')
    
    # --- Start Streaming Response --- 
    def generate():
        progress = TranslationProgress()
        translation_complete = False
        translation_error = None
        summary_text = None
        translation_thread = None
        
        # Include first_page_only in startup log
        logger.info(f"Starting process for file: {current_filename}, Generate Summary: {generate_summary}, First Page Only: {first_page_only}")

        # --- Pre-calculate Total Paragraphs for Accurate Progress --- 
        total_paragraphs_for_progress = 1 # Default to 1 to avoid division by zero
        try:
            progress.update(status="Analyzing document structure...", progress=1)
            yield progress.to_json() + '\n'
            temp_doc = Document(docx_file_to_process) # Load doc temporarily for counting
            count = len(temp_doc.paragraphs)
            for table in temp_doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        count += len(cell.paragraphs)
            for section in temp_doc.sections:
                 if section.header: count += len(section.header.paragraphs)
                 if section.footer: count += len(section.footer.paragraphs)
                 # Add others if needed (first page, even page)
            total_paragraphs_for_progress = max(1, count) # Ensure at least 1
            logger.info(f"Calculated total paragraphs for progress: {total_paragraphs_for_progress}")
            progress.update(message=f"Document has approx. {total_paragraphs_for_progress} paragraphs to process.")
            yield progress.to_json() + '\n'
            del temp_doc # Release memory
        except Exception as e:
             logger.warning(f"Could not pre-calculate total paragraphs: {e}. Progress updates might be less accurate.")
             # Keep total_paragraphs_for_progress = 1

        try: 
            # --- Define Progress Callback --- 
            def update_progress_callback(processed_count, total_count):
                # Calculate percentage (ensure total_count is not zero)
                # Start progress after initial steps (e.g., summary generation)
                base_progress = 10 if generate_summary else 5 
                # Scale translation progress to fit remaining percentage (e.g., up to 95%)
                max_translation_progress = 95
                if total_count > 0:
                    percent = base_progress + int(((processed_count / total_count) * (max_translation_progress - base_progress)))
                else: 
                    percent = base_progress # Avoid division by zero
                percent = min(max_translation_progress, percent) # Cap at max
                
                # Update the main progress object
                progress.update(progress=percent, status=f"Translating ({processed_count}/{total_count})...")
                # Note: We don't yield here directly to avoid flooding the stream
                # Instead, the main monitoring loop will yield periodically
                # Log detailed progress
                if processed_count % 50 == 0: # Log every 50 paragraphs
                     logger.debug(f"Progress callback: {processed_count}/{total_count} paragraphs processed ({percent}%)")
            
            # --- Generate Summary (if requested) --- 
            if generate_summary:
                progress.update(status="Generating summary...", progress=2, message="Reading document for summary...")
                yield progress.to_json() + '\n'
                try:
                    summary_doc = Document(docx_file_to_process)
                    full_text_list = []
                    for para in summary_doc.paragraphs:
                        full_text_list.append(para.text)
                    for table in summary_doc.tables:
                        for row in table.rows:
                            for cell in row.cells:
                                for para in cell.paragraphs:
                                    full_text_list.append(para.text)
                    full_text = "\n".join(full_text_list).strip()
                    
                    if full_text:
                         progress.update(status="Generating summary...", message="Sending text to Groq for summarization...")
                         yield progress.to_json() + '\n'
                         summary_text = summarize_text_groq(full_text, target_language=target_lang) 
                         progress.update(status="Summary generated", message="Summary received from Groq.", summary=summary_text)
                         yield progress.to_json() + '\n'
                         logger.info(f"Summary generated: {summary_text[:100]}...")
                    else:
                         progress.update(status="Skipping summary", message="Document contains no text to summarize.")
                         yield progress.to_json() + '\n'
                         logger.warning("Skipped summary generation because document text was empty.")
                except Exception as e:
                    logger.error(f"Error during summary generation: {e}")
                    progress.update(status="Summary Error", message=f"Failed to generate summary: {e}")
                    yield progress.to_json() + '\n'

            # --- Start Translation Thread --- 
            current_progress_before_translation = progress.progress # Get progress after summary step
            progress.update(status="Starting translation...", progress=max(current_progress_before_translation, 5)) # Ensure progress doesn't go backwards
            yield progress.to_json() + '\n'

            def run_translation():
                nonlocal translation_complete, translation_error
                try:
                    logger.info("Translation thread starting - passing callback")
                    # Pass the first_page_only flag to the translator
                    translate_docx_advanced(
                        docx_file_to_process, 
                        output_path,
                        engine=engine.lower(),
                        target_language=target_lang,
                        progress_callback=update_progress_callback, 
                        first_page_only=first_page_only # Pass the flag
                    )
                    
                    # --- Critical step: Copy intermediate file to final download path --- 
                    try:
                        logger.info(f"Translation function completed. Copying {output_path} to {download_path}")
                        shutil.copy2(output_path, download_path)
                        logger.info(f"Successfully copied file for download: {download_path}")
                        # Only set complete True AFTER successful copy
                        translation_complete = True 
                    except Exception as copy_err:
                        # If copy fails, it's a critical error for download
                        translation_error = f"Failed to copy translated file for download: {copy_err}"
                        logger.error(translation_error)
                        # Ensure translation_complete remains False
                        translation_complete = False 
                        
                except Exception as e:
                    # Catch errors from translate_docx_advanced itself
                    translation_error = str(e)
                    logger.error(f"Translation thread failed with exception: {str(e)}")
                    translation_complete = False # Ensure it's false on error
                    
            translation_thread = threading.Thread(target=run_translation)
            translation_thread.daemon = True
            translation_thread.start()
            
            # --- Monitor Progress (Now uses progress updated by callback) ---
            logger.info("Monitoring translation thread...")
            last_yield_time = time.time()
            while translation_thread.is_alive():
                # Yield progress periodically to update frontend, even if callback doesn't yield
                current_time = time.time()
                if current_time - last_yield_time >= 1.0: # Yield every 1 second
                     # Use the progress value updated by the callback
                     yield progress.to_json() + '\n'
                     last_yield_time = current_time
                time.sleep(0.1) # Sleep briefly
            
            # Ensure final state is yielded after thread finishes
            logger.info("Translation thread finished. Yielding final state.")
            yield progress.to_json() + '\n'
            
            # --- Final Update (Check completion/error status set by thread) --- 
            if translation_error:
                # Ensure progress reflects error state
                progress.update(status="Error", message=f"Translation failed: {translation_error}", progress=100, summary=None)
                logger.error(f"Final state: Translation failed: {translation_error}")
                yield progress.to_json() + '\n'
            elif translation_complete:
                final_message = "Document translation completed successfully!"
                # Ensure progress is 100%
                if summary_text:
                    progress.update(progress=100, status="Translation completed", message=final_message, api_calls=len(translation_cache), summary=summary_text)
                else:
                    progress.update(progress=100, status="Translation completed", message=final_message, api_calls=len(translation_cache))
                logger.info("Final state: Translation completed successfully.")
                yield progress.to_json() + '\n'
            else: # Should ideally not happen if thread completes
                progress.update(status="Unknown State", message="Translation process ended unexpectedly after thread completion.", progress=100)
                logger.warning("Final state: Translation process ended in an unknown state.")
                yield progress.to_json() + '\n'
        
        # --- Cleanup --- 
        finally:
            # Wait a bit longer and check thread status before cleaning
            join_timeout = 5 
            if translation_thread and translation_thread.is_alive():
                logger.debug(f"Waiting for translation thread to join (timeout={join_timeout}s)...")
                translation_thread.join(timeout=join_timeout) 
                if translation_thread.is_alive():
                     logger.warning("Translation thread did not join within the timeout.")
            elif translation_thread:
                 logger.debug("Translation thread already finished.")
            else:
                 logger.debug("Translation thread object not found.")
                 
            # Log state immediately before cleanup begins
            logger.info(f"Pre-cleanup state: translation_complete={translation_complete}, translation_error='{translation_error}'")
                 
            logger.debug("Starting cleanup...")
            files_to_consider = {
                'converted_docx': docx_file_to_process if original_filename.lower().endswith('.pdf') else None,
                'intermediate_translated': output_path
            }
            
            # --- Temporarily Comment Out Cleanup Loop for Debugging --- 
            # logger.info("Automatic file cleanup temporarily disabled for debugging.")
            # for name, f_path in files_to_consider.items():
            #     if f_path and os.path.exists(f_path) and f_path != download_path:
            #         try: 
            #             os.remove(f_path)
            #             logger.info(f"Removed temporary file ({name}): {f_path}")
            #         except Exception as e: 
            #             logger.error(f"Error cleaning up file {f_path} ({name}): {str(e)}")
            #     elif f_path and f_path == download_path:
            #         logger.debug(f"Skipping removal of download file: {f_path}")
            #     elif not f_path:
            #         logger.debug(f"Skipping removal for {name} as path is None")
            #     elif not os.path.exists(f_path):
            #         logger.debug(f"Skipping removal for {name} as file does not exist: {f_path}")
            # --- End of Temporarily Commented Out Block ---

    # Return the streaming response
    return Response(generate(), 
                      mimetype='text/event-stream',
                      headers={
                          'Cache-Control': 'no-cache',
                          'Connection': 'keep-alive',
                          'X-Accel-Buffering': 'no'
                      })

@app.route('/clear-cache', methods=['POST'])
def clear_cache():
    count = len(translation_cache)
    translation_cache.clear()
    return jsonify({
        'message': f'Cleared {count} cached translations.',
        'status': 'success'
    })

@app.route('/download/<filename>', methods=['GET'])
def download_file(filename):
    """Download the translated document"""
    logger.info(f"Download request received for filename from URL: {filename}")
    # Use the filename from the URL directly as the base_name
    base_name = filename 
    logger.info(f"Using base_name directly from URL: {base_name}")
    
    download_path = os.path.join(app.config['UPLOAD_FOLDER'], f'download_{base_name}.docx')
    logger.info(f"Constructed download path: {download_path}")
    
    # --- Add Directory Listing for Debugging --- 
    try:
        upload_dir_contents = os.listdir(app.config['UPLOAD_FOLDER'])
        logger.info(f"Contents of {app.config['UPLOAD_FOLDER']}: {upload_dir_contents}")
    except Exception as list_err:
        logger.error(f"Error listing directory {app.config['UPLOAD_FOLDER']}: {list_err}")
    # --- End Directory Listing --- 
    
    # Explicitly check existence and log result
    file_exists = os.path.exists(download_path)
    logger.info(f"Checking existence of {download_path}: {file_exists}")
    
    if not file_exists:
        logger.error(f"File not found at path: {download_path}")
        return "File not found or processing incomplete", 404
    
    # Set the download name (you might want to adjust this if original filenames had extensions)
    # For now, keeping the pattern consistent with potential original names
    download_name = f"translated_{base_name}.docx"
    logger.info(f"Serving file {download_path} as attachment name {download_name}")
    
    return send_file(download_path, 
                    as_attachment=True, 
                    download_name=download_name, 
                    mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

# Scheduled task to clean up old files
def cleanup_old_files():
    """Clean up files older than 30 minutes"""
    while True:
        try:
            current_time = time.time()
            for file in os.listdir(app.config['UPLOAD_FOLDER']):
                file_path = os.path.join(app.config['UPLOAD_FOLDER'], file)
                if os.path.isfile(file_path):
                    file_age = current_time - os.path.getmtime(file_path)
                    # Remove files older than 30 minutes
                    if file_age > 30 * 60:
                        os.remove(file_path)
                        logger.info(f"Removed old file: {file_path}")
        except Exception as e:
            logger.error(f"Error cleaning up old files: {e}")
        
        # Run every 10 minutes
        time.sleep(10 * 60)

# Start the cleanup thread when the app starts
cleanup_thread = threading.Thread(target=cleanup_old_files, daemon=True)
cleanup_thread.start()

if __name__ == '__main__':
    app.run(debug=True, port=5000) 