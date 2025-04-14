from docx import Document
from googletrans import Translator
import time

def translate_text(text, dest_lang='hi'):
    """
    Translate the given text to the destination language.
    This function uses googletrans. For production use, consider robust error handling.
    """
    translator = Translator()
    # Googletrans may hit rate limits, so a slight delay can help
    time.sleep(0.5)
    try:
        translated = translator.translate(text, dest=dest_lang)
        return translated.text
    except Exception as e:
        print(f"Error translating text: {text}\n{e}")
        return text  # Fallback to original if translation fails

def translate_document(source_path, target_path, dest_lang='hi'):
    # Open the source document
    doc = Document(source_path)
    # Create a new document for the translated output
    new_doc = Document()

    # Process each paragraph in the source document
    for para in doc.paragraphs:
        # Add a new paragraph in the output document
        new_para = new_doc.add_paragraph()
        # Optionally copy the style from the source paragraph
        new_para.style = para.style
        
        # Process each run to preserve formatting
        for run in para.runs:
            original_text = run.text
            if original_text.strip():
                translated_text = translate_text(original_text, dest_lang)
            else:
                translated_text = original_text
            # Add the translated run with preserved formatting
            new_run = new_para.add_run(translated_text)
            new_run.bold = run.bold
            new_run.italic = run.italic
            new_run.underline = run.underline
            # Copy font attributes if available
            if run.font.name:
                new_run.font.name = run.font.name
            if run.font.size:
                new_run.font.size = run.font.size
            if run.font.color.rgb:
                new_run.font.color.rgb = run.font.color.rgb

    # Save the translated document
    new_doc.save(target_path)
    print(f"Translation complete. Saved to {target_path}")

if __name__ == '__main__':
    # Replace 'input.docx' with the path to your converted DOCX file.
    source_file = 'input.docx'
    target_file = 'translated.docx'
    translate_document(source_file, target_file)
